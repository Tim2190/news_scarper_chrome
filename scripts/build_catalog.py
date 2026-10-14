"""Собирает каталог партии: ID | Дата | Описание (заголовок и полный текст новости с сайта + тип и оператор).

python3 scripts/build_catalog.py <черновик.xlsx | список.txt> <matches.json> data/news.json <batch.json> <каталог.xlsx>

* черновик.xlsx — колонки ID | Дата | Описание | Тип материала | Оператор | Проверить (как черновик мая 2019);
* список.txt — имена файлов по одному в строке (dir /b); тип съёмки и оператор разбираются из имени;
* batch.json — ручные уточнения партии (data/batches/*.json): exclude, notes[{match, urls, lead, own_only, check, date, keep_draft}].
"""
import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

src_path, matches_path, news_path, batch_path, out_path = sys.argv[1:6]
matches = {m["id"]: m for m in json.load(open(matches_path, encoding="utf-8"))}
news = {a["url"]: a for a in json.load(open(news_path, encoding="utf-8"))}
batch = json.load(open(batch_path, encoding="utf-8"))

# Тип съёмки и операторы по фрагментам имени файла. «<» и «>» — граница слова (подчёркивание тоже граница).
KINDS = [
    (r"krup\w*[ ._]*pl|<kr\.plan", "Крупный план"), (r"<PGM", "PGM (программный микс)"),
    (r"stacionar", "Стационар"), (r"<PTS>", "PTS (передвижная телестанция)"), (r"spec\s?rep", "Спецрепортаж"),
    (r"<sokr", "Сокращённая версия"), (r"dezhurn", "Дежурная камера"), (r"operator kanal", "Оператор канала"),
    (r"prohodka", "Проходка"), (r"lichniki", "Личники"), (r"bez elbasy", "Без Елбасы"),
    (r"<SBS", "Исходник SBS"), (r"<XDCAM", "Исходник XDCAM"),
]
OPERATORS = [
    (r"<X\.?O>|<H\.O>", "Х.О."), (r"batkeev", "Н.Баткеев"), (r"ismailov", "Исмаилов"),
    (r"tlespa[ye]{1,2}v|<D\.Tl>", "Д.Тлеспаев"), (r"kakimba[ye]{1,2}v", "Е.Какимбаев"), (r"abilov", "Ж.Абилов"),
    (r"jaitikov", "С.Жайтиков"), (r"kulmanov", "Н.Кульманов"), (r"malikov", "Маликов"),
    (r"abdrahmanov", "М.Абдрахманов"), (r"kerimkulov", "Е.Керимкулов"), (r"d(zh|j)avatarov", "Джаватаров"),
    (r"baigereev", "М.Байгереев"), (r"ahmetov|<J\.Ahm>", "Ж.Ахметов"), (r"omarkulov", "Х.Омаркулов"),
    (r"omargaliev", "Омаргалиев"), (r"<A\.?O>", "А.О."),
]


def _rx(p):
    return re.compile(p.replace("<", r"(?<![a-z])").replace(">", r"(?![a-z])"), re.I)


KINDS = [(_rx(p), k) for p, k in KINDS]
OPERATORS = [(_rx(p), o) for p, o in OPERATORS]


def kind_and_operator(name):
    kinds = [k for p, k in KINDS if p.search(name)]
    ops = [o for p, o in OPERATORS if p.search(name)]
    return "; ".join(kinds) or None, ", ".join(dict.fromkeys(ops)) or None


def load_rows():
    """[(ID, дата, описание из черновика, тип, оператор)]"""
    if src_path.lower().endswith(".xlsx"):
        ws = load_workbook(src_path).active
        for r in ws.iter_rows(min_row=2, values_only=True):
            fid, date, descr, kind, op = (list(r) + [None] * 6)[:5]
            if fid:
                yield fid, date, descr, kind, op
    else:
        for line in open(src_path, encoding="utf-8"):
            fid = line.strip()
            if fid:
                m = matches.get(fid, {})
                date = datetime.strptime(m["date"], "%d.%m.%Y") if m.get("date") else None
                yield (fid, date, None, *kind_and_operator(fid))


def site_text(url):
    """Заголовок и полный текст новости с сайта, без подписи пресс-службы и повтора заголовка."""
    a = news[url]
    title = a["title"].strip().rstrip(".")
    lines = [ln.strip() for ln in a["text"].split("\n") if ln.strip()]
    lines = [ln for ln in lines if not ln.startswith("Пресс-служба")]
    if lines and lines[0].startswith(title):
        lines[0] = lines[0][len(title):].strip()
    return title + ".\n" + "\n".join(ln for ln in lines if ln)


def describe(lead, urls, kind, operator):
    blocks = [lead.rstrip(".") + "."] if lead else []
    blocks += [site_text(u) for u in urls if u in news]
    tail = []
    if kind and not ("фотограф" in kind.lower() and lead and "фотограф" in lead.lower()):  # не повторять тип, если он уже в описании
        tail.append(kind[0].upper() + kind[1:].rstrip(".") + ".")
    if operator:
        tail.append(f"Оператор: {operator.rstrip('.')}.")
    if tail:
        blocks.append(" ".join(tail))
    return "\n\n".join(blocks)


rows, excluded = [], 0
for fid, date, draft_descr, kind, operator in load_rows():
    if any(x["match"] in fid for x in batch.get("exclude", [])):
        excluded += 1
        continue
    m = matches.get(fid, {})
    urls = [m["url"]] if m.get("url") else []
    lead, own_only, keep_draft, checks = None, False, False, []
    for n in batch.get("notes", []):
        if n["match"] not in fid:
            continue
        urls = n.get("urls", urls)
        lead = n.get("lead", lead)
        own_only = n.get("own_only", own_only)
        keep_draft = n.get("keep_draft", keep_draft)
        if n.get("date"):
            date = datetime.strptime(n["date"], "%d.%m.%Y")
        if n.get("check"):
            checks.append(n["check"])
    if keep_draft or (not urls and not lead):
        descr = draft_descr or describe(None, [], kind, operator)
    else:
        descr = describe(lead, [] if own_only else urls, kind, operator)
    conf = m.get("conf", "") if urls else "не найдено"
    if m.get("url") and urls != [m["url"]]:
        conf = "выбрано вручную"
    rows.append({
        "id": fid, "date": date, "descr": descr, "kind": kind, "operator": operator,
        "check": " ".join(checks), "urls": urls, "conf": conf,
        "title": " | ".join(news[u]["title"] for u in urls if u in news),
        "text": "\n\n".join(news[u]["text"] for u in urls if u in news),
    })

wb = Workbook()
HEAD = Font(bold=True)
FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

# Лист 1 — в формате основного каталога: ID | Дата | Описание
ws = wb.active
ws.title = "2019"
ws.append(["ID", "Дата", "Описание"])
for r in rows:
    ws.append([r["id"], r["date"], r["descr"]])
for c, w in zip("ABC", (62, 12, 110)):
    ws.column_dimensions[c].width = w

# Лист 2 — подробно, для проверки
ws2 = wb.create_sheet("Подробно")
ws2.append(["ID (имя файла в Avid)", "Дата", "Описание", "Тип материала", "Оператор / камера",
            "Проверить", "Ссылка на новость nazarbayev.kz", "Заголовок новости", "Текст новости", "Уверенность сопоставления"])
for r in rows:
    ws2.append([r["id"], r["date"], r["descr"], r["kind"], r["operator"], r["check"],
                "\n".join(r["urls"]), r["title"], r["text"], r["conf"]])
    if r["check"] or not r["urls"]:
        for cell in ws2[ws2.max_row]:
            cell.fill = FILL
for c, w in zip("ABCDEFGHIJ", (55, 12, 70, 18, 16, 45, 45, 50, 80, 16)):
    ws2.column_dimensions[c].width = w

# Лист 3 — новости сайта за месяцы, к которым относятся файлы партии
months = {(r["date"].year, r["date"].month) for r in rows if isinstance(r["date"], datetime)}
ws3 = wb.create_sheet("Новости сайта")
ws3.append(["Дата", "Рубрика", "Заголовок", "Ссылка", "Текст"])
for a in news.values():
    d = datetime.strptime(a["date"], "%d.%m.%Y") if a["date"] else None
    if not d or (d.year, d.month) not in months:
        continue
    ws3.append([d, a["rubric"], a["title"], a["url"], a["text"]])
for c, w in zip("ABCDE", (12, 14, 70, 50, 100)):
    ws3.column_dimensions[c].width = w

for sh in (ws, ws2, ws3):
    sh.freeze_panes = "A2"
    for cell in sh[1]:
        cell.font = HEAD
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            if isinstance(cell.value, datetime):
                cell.number_format = "DD.MM.YYYY"

wb.save(out_path)
print(f"Исключено: {excluded}. Файлов: {len(rows)}, со ссылкой на новость: {sum(1 for r in rows if r['urls'])}, "
      f"на проверку: {sum(1 for r in rows if r['check'])}")
